"""IMF WEO historical forecasts: what each World Economic Outlook since 1990 forecast.

Since October 2025 the WEO lives on the IMF Data Portal, and the file is linked
from the WEO dataset page there (https://data.imf.org/en/datasets/IMF.RES:WEO,
"WEO Historical Forecasts Database"). The link is read from that page, so a new
upload is found; if the page cannot be read, the last known link is tried, then
the file's older homes beside each WEO database release on imf.org:

    https://www.imf.org/-/media/files/publications/weo/weo-database/<year>/<april|october>/weohistorical.xlsx

A download that is not an Excel file (a "page not found" page served as a file,
say) counts as not found, and the next address is tried.

One Excel file from the IMF, updated after each WEO, holding every spring (April)
and fall (October) release's figures for real GDP growth, inflation and the
current account. Used for the forecast track record: how far each release's
forecasts were from what happened.

The layout read here: a sheet per variable (or variable codes in the column
names), one row per country and year, and one column per release named like
S2010 / F2010 (spring / fall 2010), optionally followed by the variable code
(S2010ngdp_rpch). The country is the ISO alpha-3 column. Anything else in the
file is ignored, and a layout this cannot read stops the fetch with a clear
error rather than storing nothing useful.

Each release is stored with breakdown `release=YYYY-MM` (04 or 10), all in one
vintage, so ordinary views (which show breakdown '') never mix them with the
current figures. Years before the release year are its estimates (E); the rest
are forecasts (P).

A copy of the file can be given instead: worldmodel-data fetch imf_weo_history --from FILE.
"""

from __future__ import annotations

import html
import logging
import re
from datetime import datetime
from pathlib import Path
from urllib.parse import urljoin

import duckdb
import httpx

from worldmodel_contract import catalogue
from worldmodel_data.connectors.base import Connector, ConnectorError

log = logging.getLogger(__name__)

DATA_PAGE = "https://data.imf.org/en/datasets/IMF.RES:WEO"
KNOWN_URL = ("https://data.imf.org/-/media/iData/External-Storage/Documents/D2AC2BB4C12C4A8CB2B6DAF1159263EB/en/"
             "WEOhistorical.xlsx")
LINK = re.compile(r"""href=["']([^"']*weohistorical\.xlsx[^"']*)["']""", re.I)
URL = "https://www.imf.org/-/media/files/publications/weo/weo-database/{year}/{month}/weohistorical.xlsx"
OLD_URL = "https://www.imf.org/-/media/Files/Publications/WEO/WEO-Database/WEOhistorical.ashx"
FILE = "WEOhistorical.xlsx"
RELEASE_COLUMN = re.compile(r"^\s*(S|F|spring|fall|apr|april|oct|october)[\s_-]?((?:19|20)\d\d)\s*([a-z_]*)\s*$", re.I)
ISO_HEADER = re.compile(r"iso.*(3|alpha)|^iso$|^iso_?code$|^country_?code$", re.I)


class IMFWEOHistory(Connector):
    source_id = "imf_weo_history"
    stage_version = 2  # 2: values outside the indicator's bounds left out

    def default_release(self, fetched_at: datetime) -> str:
        return fetched_at.strftime("%Y-%m-%d")  # the file has no release of its own: the fetch date

    def release_from_files(self, files: list[Path]) -> str | None:
        return None

    def download(self, raw_dir: Path, release: str) -> str:
        tried = []
        try:
            on = datetime.strptime(release[:10], "%Y-%m-%d")
        except ValueError:  # a release given by hand
            on = datetime.now()
        urls = [*self.linked_urls(), KNOWN_URL, *candidate_urls(on)]
        for url in dict.fromkeys(urls):
            try:
                self.get(url, raw_dir / FILE)
                if not _is_excel(raw_dir / FILE):
                    raise ConnectorError(f"GET {url} did not return an Excel file")
                return url
            except ConnectorError as exc:
                tried.append(str(exc))
                (raw_dir / FILE).unlink(missing_ok=True)
        raise ConnectorError("imf_weo_history: could not find the IMF historical forecasts file; tried:\n  "
                             + "\n  ".join(tried) + f"\nDownload it from {DATA_PAGE} (link \"WEO Historical "
                             "Forecasts Database\") and load it with: worldmodel-data fetch imf_weo_history --from FILE")

    def linked_urls(self) -> list[str]:
        """The historical forecasts link(s) on the IMF Data Portal's WEO page, if it can be read."""
        try:
            r = self.client.get(DATA_PAGE)
        except httpx.HTTPError as exc:
            log.info("imf_weo_history: could not read %s (%s); trying known addresses", DATA_PAGE, exc)
            return []
        if r.status_code != 200:
            log.info("imf_weo_history: %s returned HTTP %s; trying known addresses", DATA_PAGE, r.status_code)
            return []
        return [urljoin(DATA_PAGE, html.unescape(m)) for m in LINK.findall(r.text)]

    def stage(self, con: duckdb.DuckDBPyConnection, raw_dir: Path, release: str, vintage: str) -> str:
        files = sorted(p for p in raw_dir.iterdir() if p.suffix.lower() in (".xlsx", ".xlsm", ".ashx"))
        if not files:
            raise ConnectorError(f"imf_weo_history: no Excel file in {raw_dir}")
        rows = read_workbook(files[0], self.series_map())
        if not rows:
            raise ConnectorError(f"imf_weo_history: found no release columns (like S2010, F2010) in {files[0].name}; "
                                 "the IMF may have changed the file's layout")
        rows = within_bounds(rows)
        con.execute("CREATE TABLE weo_hist (source_code VARCHAR, indicator_id VARCHAR, period_start DATE, "
                    "breakdown VARCHAR, value DOUBLE, status VARCHAR)")
        con.executemany("INSERT INTO weo_hist VALUES (?, ?, ?, ?, ?, ?)", rows)
        return (f"SELECT 'imf_weo_history' AS source_id, '{vintage}' AS vintage, 'imf' AS scheme, source_code, "
                f"indicator_id, period_start, 'A' AS freq, breakdown, 'central' AS variant, value, status "
                f"FROM weo_hist")


def within_bounds(rows: list[tuple]) -> list[tuple]:
    """Leave out values outside the catalogue's bounds for their indicator. The file keeps
    every number each release printed, slips included (São Tomé's 2012 growth once
    appears as 297%), and one such value would stop the whole build."""
    bounds = {i.id: i.bounds for i in catalogue.indicators().values()}
    kept, dropped = [], []
    for r in rows:
        lo, hi = bounds.get(r[1], (float("-inf"), float("inf")))
        (kept if lo <= r[4] <= hi else dropped).append(r)
    for r in dropped[:20]:
        log.warning("imf_weo_history: left out %s %s %s (%s): %s is outside %s", r[0], r[1], r[2][:4], r[3], r[4],
                    bounds[r[1]])
    if len(dropped) > 20:
        log.warning("imf_weo_history: and %s more values outside their bounds", len(dropped) - 20)
    return kept


def candidate_urls(on: datetime) -> list[str]:
    """Where the file may be: beside the latest WEO release as of `on` and the four before it, then the old address."""
    year, month = on.year, on.month
    releases = []
    # WEO comes out in mid April and mid October; count back from the latest one out by now.
    y, m = (year, 10) if (month, on.day) >= (10, 15) else (year, 4) if (month, on.day) >= (4, 15) else (year - 1, 10)
    for _ in range(5):
        releases.append((y, m))
        y, m = (y, 4) if m == 10 else (y - 1, 10)
    return [URL.format(year=y, month="october" if m == 10 else "april") for y, m in releases] + [OLD_URL]


def _is_excel(path: Path) -> bool:
    """An .xlsx file is a zip; a web page saved under that name is not."""
    with path.open("rb") as fh:
        return fh.read(4) == b"PK\x03\x04"


def read_workbook(path: Path, series: dict[str, tuple[str, float]]) -> list[tuple]:
    """(country, indicator, period_start, breakdown, value, status) for every release column found."""
    from openpyxl import load_workbook  # only needed here

    codes = {c.lower(): v for c, v in series.items()}
    fh = path.open("rb")  # a file object, so a copy saved as .ashx opens too
    wb = load_workbook(fh, read_only=True, data_only=True)
    out: list[tuple] = []
    try:
        for ws in wb.worksheets:
            sheet_code = next((c for c in codes if c in ws.title.lower()), None)
            it = ws.iter_rows(values_only=True)
            header, iso_col, year_col, releases = None, None, None, {}
            for _ in range(15):  # the header is near the top
                row = next(it, None)
                if row is None:
                    break
                cells = [str(c).strip() if c is not None else "" for c in row]
                found = {}
                for i, cell in enumerate(cells):
                    m = RELEASE_COLUMN.match(cell)
                    if not m:
                        continue
                    season, year, suffix = m.group(1).lower(), m.group(2), m.group(3).lower()
                    code = next((c for c in codes if suffix and c in suffix), None) if suffix else sheet_code
                    if code:
                        found[i] = (f"{year}-{'04' if season[0] in 'sa' else '10'}", code)
                if found:
                    header, releases = cells, found
                    iso_col = next((i for i, c in enumerate(cells) if ISO_HEADER.search(c)), None)
                    year_col = next((i for i, c in enumerate(cells) if c.lower() == "year"), None)
                    break
            if header is None or iso_col is None or year_col is None:
                if header is not None:
                    log.warning("imf_weo_history: sheet %s has release columns but no ISO code or year column", ws.title)
                continue
            for row in it:
                if row is None or len(row) <= max(iso_col, year_col):
                    continue
                iso, year = row[iso_col], row[year_col]
                try:
                    year = int(year)
                except (TypeError, ValueError):
                    continue
                if not iso or not isinstance(iso, str):
                    continue
                for i, (rel, code) in releases.items():
                    if i >= len(row):
                        continue
                    value = _number(row[i])
                    if value is None:
                        continue
                    indicator, scale = codes[code]
                    out.append((iso.strip().upper(), indicator, f"{year}-01-01", f"release={rel}", value * scale,
                                "P" if year >= int(rel[:4]) else "E"))
    finally:
        wb.close()
        fh.close()
    return out


def _number(v) -> float | None:
    if isinstance(v, (int, float)):
        return float(v)
    if isinstance(v, str):
        s = v.strip().replace(",", "")
        if s in ("", "n/a", "--", "NA", ".", "..."):
            return None
        try:
            return float(s)
        except ValueError:
            return None
    return None
