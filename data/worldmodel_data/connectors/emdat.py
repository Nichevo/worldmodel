"""Disasters since 1900: EM-DAT, the international disaster database (CRED, UCLouvain).

    https://public.emdat.be/  ->  sign in (free for non-commercial use), Data,
    tick "Natural" (and "Technological" if wanted), Download  ->  an .xlsx file

The download needs an account, so a refresh never fetches it: load the file with
    worldmodel-data fetch emdat --from <the .xlsx>
and do the same again when you want newer data (once a year is plenty).

One row per disaster and country: DisNo., Disaster Type and Subtype, Event Name,
ISO, Country, Location, Start/End Year/Month/Day, Total Deaths, Total Affected and
Total Damage, Adjusted ('000 US$, in today's dollars), Magnitude and Magnitude
Scale. Headers are matched loosely (case, punctuation and the odd renamed column
do not matter), and a CSV export works too. Each row becomes a disaster event;
its size is the deaths, or else the people affected. Epidemics are their own kind.
"""

from __future__ import annotations

import csv
import io
import logging
import re
from datetime import datetime
from pathlib import Path

import duckdb

from worldmodel_data.connectors.base import Connector, ConnectorError
from worldmodel_data.connectors.events_common import EMPTY_OBSERVATIONS, country, make_date, stage_events

log = logging.getLogger(__name__)

# Field -> header names it may go by, compared after lower-casing and dropping all but letters and digits.
FIELDS = {
    "id": ("disno", "disasterno", "disno."),
    "group": ("disastergroup",),
    "type": ("disastertype",),
    "subtype": ("disastersubtype",),
    "name": ("eventname",),
    "iso": ("iso", "iso3", "isocode"),
    "country": ("country",),
    "location": ("location",),
    "start_year": ("startyear", "year"),
    "start_month": ("startmonth",),
    "start_day": ("startday",),
    "end_year": ("endyear",),
    "end_month": ("endmonth",),
    "end_day": ("endday",),
    "deaths": ("totaldeaths",),
    "affected": ("totalaffected",),
    "damage_adj": ("totaldamageadjusted000us", "totaldamageadjusted000usd"),
    "damage": ("totaldamage000us", "totaldamage000usd"),
    "magnitude": ("magnitude", "dismagvalue"),
    "scale": ("magnitudescale", "dismagscale"),
}


def _key(header: str) -> str:
    return re.sub(r"[^a-z0-9]", "", str(header or "").lower())


def columns(header: list) -> dict[str, int]:
    keys = [_key(h) for h in header]
    out = {}
    for field, names in FIELDS.items():
        wanted = {_key(n) for n in names}
        i = next((i for i, k in enumerate(keys) if k in wanted), None)
        if i is not None:
            out[field] = i
    return out


class EMDAT(Connector):
    source_id = "emdat"
    manual = True

    def default_release(self, fetched_at: datetime) -> str:
        return fetched_at.strftime("%Y-%m-%d")

    def release_from_files(self, files: list[Path]) -> str | None:
        for f in files:  # the public table's file name carries its date: public_emdat_custom_request_2025-06-01_...
            m = re.search(r"(\d{4}-\d{2}-\d{2})", f.name)
            if m:
                return m.group(1)
        return None

    def download(self, raw_dir: Path, release: str) -> str:
        raise ConnectorError("emdat: EM-DAT needs a (free) sign-in, so it cannot be downloaded for you. Download the "
                             "table from https://public.emdat.be/ and load it with: "
                             "worldmodel-data fetch emdat --from <file.xlsx>")

    def stage(self, con: duckdb.DuckDBPyConnection, raw_dir: Path, release: str, vintage: str) -> str:
        files = sorted(p for p in raw_dir.iterdir() if p.suffix.lower() in (".xlsx", ".csv"))
        if not files:
            raise ConnectorError(f"emdat: no .xlsx or .csv file in {raw_dir}")
        rows, unmatched = read(table(files[0]))
        if not rows:
            raise ConnectorError(f"emdat: no disasters read from {files[0].name}; is it the EM-DAT table?")
        log.info("emdat: %s disaster rows", len(rows))
        self.events_query = stage_events(con, self.source_id, vintage, rows, unmatched)
        return EMPTY_OBSERVATIONS


def table(path: Path) -> list[list]:
    """All rows of the first sheet (or the CSV), header first."""
    if path.suffix.lower() == ".csv":
        return list(csv.reader(io.StringIO(path.read_text(encoding="utf-8-sig", errors="replace"))))
    from openpyxl import load_workbook  # only needed here

    with path.open("rb") as fh:
        wb = load_workbook(fh, read_only=True, data_only=True)
        try:
            return [list(r) for r in wb.worksheets[0].iter_rows(values_only=True)]
        finally:
            wb.close()


def _num(v) -> float | None:
    if v is None or v == "":
        return None
    try:
        return float(str(v).replace(",", ""))
    except ValueError:
        return None


def read(rows: list[list]) -> tuple[list[tuple], set[str]]:
    # The header is the first row naming a disaster type and a start year.
    start = next((i for i, r in enumerate(rows[:20]) if {"type", "start_year"} <= set(columns(r))), None)
    if start is None:
        return [], set()
    col = columns(rows[start])
    known = {c for c in col}
    get = lambda r, f: (r[col[f]] if f in known and col[f] < len(r) else None)  # noqa: E731
    out, unmatched, seen = [], set(), {}
    for r in rows[start + 1:]:
        if not r or not any(v not in (None, "") for v in r):
            continue
        begin, prec = make_date(get(r, "start_year"), get(r, "start_month"), get(r, "start_day"))
        if begin is None:
            continue
        end, _ = make_date(get(r, "end_year"), get(r, "end_month"), get(r, "end_day"))
        iso = str(get(r, "iso") or "").strip().upper()
        name = str(get(r, "country") or "").strip()
        entity = iso if re.fullmatch(r"[A-Z]{3}", iso) and _known(iso) else country(name, begin, unmatched)
        if entity is None:
            continue
        dtype = str(get(r, "type") or "").strip() or "Disaster"
        sub = str(get(r, "subtype") or "").strip() or None
        kind = "epidemic" if dtype.lower() == "epidemic" else "disaster"
        deaths, affected = _num(get(r, "deaths")), _num(get(r, "affected"))
        damage_k = _num(get(r, "damage_adj"))
        if damage_k is None:
            damage_k = _num(get(r, "damage"))
        mag, scale = _num(get(r, "magnitude")), str(get(r, "scale") or "").strip()
        title = str(get(r, "name") or "").strip() or (f"{sub} ({dtype.lower()})" if sub and sub != dtype else dtype)
        dis_no = str(get(r, "id") or "").strip()
        n = seen[(dis_no, entity)] = seen.get((dis_no, entity), 0) + 1
        key = f"{dis_no or begin.isoformat() + ':' + dtype}:{entity}" + (f":{n}" if n > 1 else "")
        size, unit = (deaths, "deaths") if deaths else (affected, "people affected") if affected else (None, None)
        detail = "; ".join(x for x in (
            f"{dtype}{' / ' + sub if sub and sub != dtype else ''}",
            f"magnitude {mag:g} {scale}".strip() if mag is not None else "",
            f"EM-DAT {dis_no}" if dis_no else "",
        ) if x)
        out.append((entity, key, kind, (sub or dtype).lower(), title, begin, end if end and end >= begin else None,
                    prec, size, unit, deaths, affected, damage_k * 1000 if damage_k is not None else None,
                    str(get(r, "location") or "").strip()[:300] or None, None, detail))
    return out, unmatched


_ISO3: set[str] | None = None


def _known(iso: str) -> bool:
    global _ISO3
    if _ISO3 is None:
        from worldmodel_contract import catalogue
        _ISO3 = {c["code"] for c in catalogue.entity_codes() if c["scheme"] == "iso3"}
    return iso in _ISO3
