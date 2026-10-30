"""Shared helpers for connectors that read a table of values per country and year
and list their own indicators: collect rows in Python, stage them in one go."""

from __future__ import annotations

import logging
import re
from collections import defaultdict
from datetime import date

import duckdb

from worldmodel_contract import catalogue
from worldmodel_data.connectors.base import sql_str

log = logging.getLogger(__name__)


class Rows:
    """Values by (scheme, code, indicator, year, breakdown), with a status each.
    Quarterly or monthly values added with `add_period` are averaged to the year."""

    def __init__(self) -> None:
        self.rows: list[tuple] = []
        self._periods: dict[tuple, list[float]] = defaultdict(list)

    def add(self, scheme: str, code: str, indicator: str, year: int, value: float | None,
            breakdown: str = "", status: str = "E") -> None:
        if value is None or value != value or abs(value) == float("inf"):
            return
        self.rows.append((scheme, str(code), indicator, date(int(year), 1, 1), breakdown, float(value), status))

    def add_period(self, scheme: str, code: str, indicator: str, year: int, value: float | None,
                   breakdown: str = "") -> None:
        if value is None or value != value:
            return
        self._periods[(scheme, str(code), indicator, int(year), breakdown)].append(float(value))

    def finish_periods(self, need: int = 1) -> None:
        """Turn the values added by period into yearly means, for years with at least `need` of them."""
        for (scheme, code, ind, year, bd), vals in self._periods.items():
            if len(vals) >= need:
                self.add(scheme, code, ind, year, sum(vals) / len(vals), bd)
        self._periods.clear()

    def __len__(self) -> int:
        return len(self.rows)

    def stage(self, con: duckdb.DuckDBPyConnection, source_id: str, vintage: str, table: str = "vals") -> str:
        con.execute(f"CREATE TABLE {table} (scheme VARCHAR, source_code VARCHAR, indicator_id VARCHAR, "
                    "period_start DATE, breakdown VARCHAR, value DOUBLE, status VARCHAR)")
        if self.rows:
            con.executemany(f"INSERT INTO {table} VALUES (?, ?, ?, ?, ?, ?, ?)", self.rows)
        return (f"SELECT {sql_str(source_id)} AS source_id, {sql_str(vintage)} AS vintage, scheme, source_code, "
                f"indicator_id, period_start, 'A' AS freq, breakdown, 'central' AS variant, value, status FROM {table}")


def indicator(id: str, name: str, description: str, unit: str, domain: str, source_id: str, series: str, *,
              decimals: int = 1, aggregation: dict | None = None, default_on: bool = True, notes: str = "",
              breakdowns: list[str] | None = None, bounds: tuple[float, float] = (-1e18, 1e18)) -> dict:
    """An indicator a source lists itself (indicators.json), catalogue-shaped."""
    return {
        "id": id, "name": name, "description": description, "domain": domain, "unit": unit, "decimals": decimals,
        "aggregation": aggregation or {"method": "none"}, "bounds": list(bounds), "breakdowns": breakdowns or [],
        "sources": [{"source_id": source_id, "series": series, "scale": 1.0, "notes": notes}],
        "default_on": default_on,
    }


def num(v) -> float | None:
    """A number from a cell: blanks, dots, dashes and footnote marks give None; "1,234" and "56%" work."""
    if v is None:
        return None
    if isinstance(v, (int, float)):
        return float(v)
    s = str(v).strip().replace(",", "").replace("%", "").replace(" ", "")
    if s in ("", "..", ".", "-", "–", "—", "n/a", "NA", "N/A"):
        return None
    try:
        return float(s)
    except ValueError:
        m = re.fullmatch(r"(-?\d+(?:\.\d+)?)\s*\D{0,3}", s)  # "12.3 a", "45*"
        return float(m.group(1)) if m else None


def key(header) -> str:
    """A header compared loosely: lower case, letters and digits only."""
    return re.sub(r"[^a-z0-9]", "", str(header or "").lower())


_CODES: dict[str, dict[str, str]] | None = None


def entity_for(scheme: str, code) -> str | None:
    """The entity id for a code in one of the crosswalk's schemes (iso3, iso2, m49 ...)."""
    global _CODES
    if _CODES is None:
        _CODES = defaultdict(dict)
        for c in catalogue.entity_codes():
            if not c.get("valid_to"):
                _CODES[c["scheme"]][c["code"]] = c["entity_id"]
            else:
                _CODES[c["scheme"]].setdefault(c["code"], c["entity_id"])
    if code is None:
        return None
    s = str(code).strip()
    if scheme == "m49":  # UN M49 codes: 004 for Afghanistan; regions (900 the World) from the UN WPP list
        n = s.split(".")[0].lstrip("0") or "0"
        return _CODES["m49"].get(n.zfill(3)) or _CODES["un_wpp"].get(n)
    return _CODES.get(scheme, {}).get(s)


def xlsx_rows(path, sheet: int | str = 0) -> list[list]:
    """All rows of one sheet of an Excel file (by position or name)."""
    from openpyxl import load_workbook  # only needed here

    with open(path, "rb") as fh:
        wb = load_workbook(fh, read_only=True, data_only=True)
        try:
            ws = wb.worksheets[sheet] if isinstance(sheet, int) else wb[sheet]
            return [list(r) for r in ws.iter_rows(values_only=True)]
        finally:
            wb.close()


def xlsx_sheet_names(path) -> list[str]:
    from openpyxl import load_workbook

    with open(path, "rb") as fh:
        wb = load_workbook(fh, read_only=True)
        try:
            return list(wb.sheetnames)
        finally:
            wb.close()


def is_country(entity_id: str | None) -> bool:
    """True for a country or former state (not a region or group)."""
    e = catalogue.entities().get(entity_id or "")
    return bool(e) and not e.is_group
